"""Phase 2: money you can trust.

  * integer paisa and the one rounding rule (to_paisa / to_rupees)
  * reversals instead of edits, and append-only history enforced by the database
  * stock can never go below zero, on any path
  * every on_hand change is in the stock ledger
  * gapless document numbers
  * the V5 migration of a real pre-V5 (REAL-float) business file, and that it cannot half-apply
"""
from __future__ import annotations

import json
import random
import sqlite3
from decimal import Decimal

import pytest

from munshi.domain import migrations
from munshi.domain.models import Product, discounted_paisa, mul_div, to_business_date, to_paisa, to_rupees, today_iso
from munshi.domain.repository import InsufficientStockError, MunshiRepository, NotFoundError, StateError
from munshi.domain.seed import seeded_repository


@pytest.fixture
def repo():
    return seeded_repository()


def _year() -> str:
    return today_iso()[:4]


# ================================================================ rounding rule
@pytest.mark.parametrize("rupees, paisa", [
    (0, 0), (1, 100), (3850, 385000), (-5, -500),
    (0.1, 10), (0.1 + 0.2, 30),                  # 0.30000000000000004 -> 30, not 30.000000000000004
    (2.675, 268), (-2.675, -268),                # half a paisa rounds away from zero, symmetrically
    (1.005, 101), (0.004, 0), (0.005, 1), (-0.005, -1),
    ("12.345", 1235), (Decimal("99.995"), 10000), ("  7.5 ", 750),
    (3773.0000000000005, 377300),               # legacy float noise is absorbed
])
def test_to_paisa_rounds_half_up_on_the_written_decimal(rupees, paisa):
    assert to_paisa(rupees) == paisa and isinstance(to_paisa(rupees), int)


@pytest.mark.parametrize("bad", [None, True, False, float("nan"), float("inf"), "abc", "", 10 ** 13, -(10 ** 13)])
def test_to_paisa_refuses_what_is_not_an_amount(bad):
    with pytest.raises(ValueError):
        to_paisa(bad)


def test_to_rupees_and_the_round_trip_are_exact():
    assert to_rupees(12345) == 123.45 and to_rupees(-1) == -0.01 and to_rupees(0) == 0.0
    with pytest.raises(TypeError):
        to_rupees(1.5)
    with pytest.raises(TypeError):
        to_rupees(True)
    rnd = random.Random(7)
    for _ in range(20000):                       # every paisa amount survives rupees -> JSON float -> paisa
        p = rnd.randint(-(10 ** 14), 10 ** 14)
        assert to_paisa(to_rupees(p)) == p
        assert to_paisa(json.loads(json.dumps(to_rupees(p)))) == p


def test_mul_div_and_discount_use_the_same_rule():
    assert mul_div(10, 1, 4) == 3 and mul_div(10, 1, 3) == 3 and mul_div(-10, 1, 4) == -3 and mul_div(5, 1, 2) == 3
    assert mul_div(110011000, 30, 300) == 11001100
    assert discounted_paisa(385000, 2) == 377300 and discounted_paisa(385000, 2.5) == 375375 and discounted_paisa(625000, 2.5) == 609375
    assert discounted_paisa(333, 50) == 167          # 166.5 -> 167


def test_float_sum_regression_balance_is_exactly_zero(repo):
    """The reviewers' -5.55e-17: a khata that should be zero summed to a float residue."""
    repo.add_ledger("C-006", "invoice", 0.3, "x", None, "t")
    repo.record_payment("C-006", 0.1, "cash", "", "t")
    repo.record_payment("C-006", 0.2, "cash", "", "t")
    assert repo.outstanding_paisa("C-006") == 0 and repo.outstanding("C-006") == 0.0
    assert repo._one("SELECT typeof(SUM(amount)) t FROM ledger WHERE customer_id='C-006'")["t"] == "integer"


def test_money_columns_are_integer_paisa_and_refuse_fractions(repo):
    for table, col in (("ledger", "amount"), ("supplier_ledger", "amount"), ("expenses", "amount"), ("products", "unit_price"),
                       ("products", "cost_price"), ("customers", "credit_limit"), ("purchases", "total"), ("stops", "cash_collected"),
                       ("deposits", "amount_counted"), ("reminders", "amount_due"), ("promises", "amount")):
        decl = {r["name"]: r["type"] for r in repo._all(f"PRAGMA table_info({table})")}[col]
        assert decl == "INTEGER", (table, col, decl)
        assert repo._one(f"SELECT COUNT(*) n FROM {table} WHERE typeof({col}) != 'integer'")["n"] == 0, (table, col)
    with pytest.raises(sqlite3.IntegrityError):     # a rupee float with a fraction can't slip into a paisa column
        with repo._tx() as c:
            c.execute("INSERT INTO ledger (entry_id, customer_id, kind, amount, ref, created_at) VALUES ('X-1','C-001','invoice',1234.5,'r','2026-01-01')")


def test_api_surface_stays_in_rupees(repo):
    e = repo.record_payment("C-001", 1234.56, "bank", "tx", "t")
    assert e.amount == -1234.56 and repo._one("SELECT amount FROM ledger WHERE entry_id=?", (e.entry_id,))["amount"] == -123456
    assert repo.get_product("UREA-50").unit_price == 3850.0 and repo.get_customer("C-001").credit_limit == 1_200_000.0


# ================================================================ append-only history
@pytest.mark.parametrize("table", ["ledger", "supplier_ledger", "audit", "stock_moves", "expenses"])
def test_history_tables_are_append_only_even_for_direct_sql(repo, table):
    repo.audit("t", "x", "y", "z", {})
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        with repo._tx() as c:
            c.execute(f"UPDATE {table} SET created_at = created_at")
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        with repo._tx() as c:
            c.execute(f"DELETE FROM {table}")
    assert repo._one(f"SELECT COUNT(*) n FROM {table}")["n"] > 0


# ================================================================ reversals
def test_bounced_cheque_is_a_reversal_and_nets_everything(repo):
    before = repo.outstanding_paisa("C-005")
    repo.log_promise("C-005", 100000, today_iso(), "w")
    pay = repo.record_payment("C-005", 100000, "cheque", "chq 1142", "hisaab_munshi", "clerk")
    assert repo.open_promise("C-005")["kept"] is True
    rev = repo.reverse_ledger_entry(pay.entry_id, "cheque bounced", "owner", "owner")
    assert rev.reversal_of == pay.entry_id and rev.amount == 100000.0 and rev.kind == "payment"
    assert rev.entry_id == rev.doc_no == f"REV-{_year()}-000001"
    assert repo.outstanding_paisa("C-005") == before                       # back to where it was
    assert repo.get_ledger_entry(pay.entry_id).amount == -100000.0         # the original is untouched history
    assert repo.open_promise("C-005")["kept"] is False                     # a bounced cheque doesn't keep a promise
    row = next(r for r in repo.aging() if r["customer_id"] == "C-005")
    assert row["balance"] == to_rupees(before)
    col = repo.collection_report(today_iso(), today_iso())
    assert col["collected"] == 0.0 and col["by_method"]["cheque"] == 0.0
    with pytest.raises(StateError, match="already reversed"):
        repo.reverse_ledger_entry(pay.entry_id, "again", "owner", "owner")
    with pytest.raises(StateError, match="itself a reversal"):
        repo.reverse_ledger_entry(rev.entry_id, "undo", "owner", "owner")
    with pytest.raises(StateError, match="owner"):
        repo.reverse_ledger_entry("INV-SEED00", "no approver", "clerk", "")
    with pytest.raises(ValueError, match="reason"):
        repo.reverse_ledger_entry("INV-SEED00", "", "owner", "owner")
    assert any(a["action"] == "reverse_ledger_entry" for a in repo.audit_log(20, pay.entry_id))


def test_cash_receipt_reversal_shows_as_negative_cash_in_on_the_day_it_is_posted(repo):
    pay = repo.record_payment("C-001", 5000, "cash", "", "h", received_by="office")
    repo.reverse_ledger_entry(pay.entry_id, "keyed to the wrong customer", "owner", "owner")
    cb = repo.cashbook(today_iso())
    assert [i["amount"] for i in cb["cash_in"]] == [5000.0, -5000.0] and cb["total_in"] == 0.0


def test_expense_reversal(repo):
    today = today_iso()
    base = repo.expenses_paisa(today, today)
    x = repo.record_expense("fuel", 55000, "diesel (typo)", "cash", "Bilal", "h")
    r = repo.reverse_expense(x.expense_id, "should have been 5,500", "owner", "owner")
    assert r.amount == -55000.0 and r.reversal_of == x.expense_id and r.category == "fuel"
    repo.record_expense("fuel", 5500, "diesel", "cash", "Bilal", "h")
    assert repo.expenses_paisa(today, today) == base + 550000
    assert repo.profit_summary(today, today)["expenses"] == to_rupees(base + 550000)
    with pytest.raises(StateError):
        repo.reverse_expense(x.expense_id, "again", "owner", "owner")


def test_purchase_reversal_takes_goods_and_bill_back(repo):
    s0, bal0 = repo.get_stock("WH-VEHARI", "SEED-WHEAT").on_hand, repo.supplier_balance_paisa("S-002")
    pool0 = repo._pool("SEED-WHEAT")[1]
    pur = repo.record_purchase("S-002", "WH-VEHARI", [{"sku": "SEED-WHEAT", "qty": 40, "unit_cost": 4800}], "ENG-1", 50000, "k")
    assert pur.purchase_id == pur.doc_no == f"PUR-{_year()}-000001"
    ret = repo.reverse_purchase(pur.purchase_id, "wrong supplier invoice", "owner", "owner")
    assert ret.purchase_id == f"PRN-{_year()}-000001" and ret.reversal_of == pur.purchase_id and ret.total == -192000.0
    assert repo.get_stock("WH-VEHARI", "SEED-WHEAT").on_hand == s0 and repo.supplier_balance_paisa("S-002") == bal0
    assert repo._pool("SEED-WHEAT")[1] == pool0
    with pytest.raises(StateError):
        repo.reverse_purchase(pur.purchase_id, "again", "owner", "owner")


def test_purchase_reversal_is_refused_once_the_goods_are_gone_and_nothing_changes(repo):
    pur = repo.record_purchase("S-002", "WH-VEHARI", [{"sku": "SEED-WHEAT", "qty": 40, "unit_cost": 4800}], "ENG-1", 0, "k")
    on_hand = repo.get_stock("WH-VEHARI", "SEED-WHEAT").on_hand
    repo.adjust_stock("WH-VEHARI", "SEED-WHEAT", -(on_hand - 10), "sold off the books", "owner", "owner")
    moves, bal = len(repo.stock_moves(limit=10000)), repo.supplier_balance_paisa("S-002")
    with pytest.raises(InsufficientStockError, match="below zero"):
        repo.reverse_purchase(pur.purchase_id, "too late", "owner", "owner")
    assert repo.get_stock("WH-VEHARI", "SEED-WHEAT").on_hand == 10 and len(repo.stock_moves(limit=10000)) == moves
    assert repo.supplier_balance_paisa("S-002") == bal and repo._one("SELECT COUNT(*) n FROM purchases WHERE reversal_of IS NOT NULL")["n"] == 0
    assert repo._one("SELECT last_value FROM document_counters WHERE kind='purchase_return'") is None      # number not burnt


def test_supplier_payment_reversal_and_purchase_bill_guard(repo):
    e = repo.pay_supplier("S-001", 540000, "cheque", "chq 9", "k", "owner")
    assert repo.payables() == []
    r = repo.reverse_supplier_entry(e.entry_id, "cheque bounced", "owner", "owner")
    assert r.reversal_of == e.entry_id and repo.supplier_balance("S-001") == 540000.0
    pur = repo.record_purchase("S-003", "WH-MULTAN", [{"sku": "CYPER-1L", "qty": 5, "unit_cost": 1200}], "", 0, "k")
    bill = next(x for x in repo.supplier_ledger("S-003") if x.ref == pur.purchase_id)
    with pytest.raises(StateError, match="reverse the purchase"):
        repo.reverse_supplier_entry(bill.entry_id, "wrong", "owner", "owner")


# ================================================================ stock: never below zero, always in the ledger
def _stock_state(repo):
    return ([tuple(r) for r in repo._all("SELECT * FROM stock ORDER BY warehouse_id, sku")],
            repo._one("SELECT COUNT(*) n FROM stock_moves")["n"],
            [tuple(r) for r in repo._all("SELECT * FROM inventory_value ORDER BY sku")])


def test_adjust_below_zero_is_refused_cleanly(repo):
    before = _stock_state(repo)
    with pytest.raises(InsufficientStockError, match="never go below zero"):
        repo.adjust_stock("WH-MULTAN", "CYPER-1L", -9, "damaged", "owner", "owner")     # 8 on hand
    assert _stock_state(repo) == before
    repo.adjust_stock("WH-MULTAN", "CYPER-1L", -8, "damaged", "owner", "owner")          # exactly to zero is fine
    assert repo.get_stock("WH-MULTAN", "CYPER-1L").on_hand == 0 and repo._pool("CYPER-1L") == (60, 60 * 125000)   # Vehari's 60 keep their cost
    with pytest.raises(ValueError):
        repo.adjust_stock("WH-MULTAN", "CYPER-1L", 0, "nothing", "owner", "owner")


def test_transfer_and_allocation_beyond_stock_are_refused(repo):
    before = _stock_state(repo)
    with pytest.raises(InsufficientStockError):
        repo.transfer_stock("WH-MULTAN", "WH-VEHARI", "CYPER-1L", 9, "g", "c")
    o = repo.create_order("C-003", [{"sku": "CYPER-1L", "qty": 9}], "t", "", "o")
    repo.confirm_order(o.order_id, "o", "c")
    with pytest.raises(InsufficientStockError):
        repo.allocate_order(o.order_id, "WH-MULTAN", "g")
    # an order that splits one product over two lines is one demand, not two that each fit
    o2 = repo.create_order("C-003", [{"sku": "CYPER-1L", "qty": 5}, {"sku": "CYPER-1L", "qty": 5}], "t", "", "o")
    repo.confirm_order(o2.order_id, "o", "c")
    with pytest.raises(InsufficientStockError):
        repo.allocate_order(o2.order_id, "WH-MULTAN", "g")
    assert _stock_state(repo) == before and repo.get_order(o2.order_id).status == "confirmed"


def test_loading_a_plan_that_would_go_below_zero_moves_nothing(repo):
    o = repo.create_order("C-003", [{"sku": "UREA-50", "qty": 5}, {"sku": "CYPER-1L", "qty": 6}], "t", "", "o")   # UREA moves first, then CYPER fails
    repo.confirm_order(o.order_id, "o", "c"); repo.allocate_order(o.order_id, "WH-MULTAN", "g")
    p = repo.create_dispatch_plan(today_iso(), "R-MULTAN-N", "V-01", [o.order_id], "g")
    repo.adjust_stock("WH-MULTAN", "CYPER-1L", -4, "damaged in the godown", "owner", "owner")   # 4 left, 6 reserved
    before = _stock_state(repo)
    with pytest.raises(InsufficientStockError, match="CYPER-1L"):
        repo.approve_dispatch_plan(p.plan_id, "g", "c")
    assert _stock_state(repo) == before                        # UREA was not taken out either: all or nothing
    assert repo.get_plan(p.plan_id).status == "planned" and repo.get_order(o.order_id).status == "allocated"


def test_set_stock_is_ledgered_and_never_negative(repo):
    repo.set_stock("WH-MULTAN", "ZINC-10", 150)
    m = repo.stock_moves("ZINC-10", "WH-MULTAN")[0]
    assert (m.kind, m.delta, m.ref) == ("adjust", 10, "set_stock")
    with pytest.raises(InsufficientStockError):
        repo.set_stock("WH-MULTAN", "ZINC-10", -1)
    replay = repo.replay_stock_ledger()
    assert all(replay.get((s.warehouse_id, s.sku), 0) == s.on_hand for s in repo.list_stock())


def test_database_itself_refuses_negative_stock(repo):
    with pytest.raises(sqlite3.DatabaseError, match="below zero"):
        with repo._tx() as c:
            c.execute("UPDATE stock SET on_hand = -1 WHERE warehouse_id='WH-MULTAN' AND sku='UREA-50'")
    with pytest.raises(sqlite3.DatabaseError, match="below zero"):
        with repo._tx() as c:
            c.execute("INSERT INTO stock VALUES ('WH-MULTAN', 'NEW-SKU', -5, 0)")


def test_ledger_replays_to_stock_after_every_kind_of_move(repo):
    o = repo.create_order("C-002", [{"sku": "UREA-50", "qty": 10}, {"sku": "DAP-50", "qty": 2}], "t", "", "o")
    repo.confirm_order(o.order_id, "o", "c"); repo.allocate_order(o.order_id, "WH-MULTAN", "g")
    p = repo.create_dispatch_plan(today_iso(), "R-MULTAN-N", "V-01", [o.order_id], "g"); repo.approve_dispatch_plan(p.plan_id, "g", "c")
    st = repo.list_stops(p.plan_id)[0]
    repo.close_stop(st.stop_id, [{"sku": "UREA-50", "qty": 7}], [{"sku": "UREA-50", "qty": 2}, {"sku": "DAP-50", "qty": 2}], 0, st.otp, "d")
    repo.record_purchase("S-001", "WH-MULTAN", [{"sku": "DAP-50", "qty": 10, "unit_cost": 6000}], "", 0, "k")
    repo.transfer_stock("WH-MULTAN", "WH-VEHARI", "DAP-50", 5, "g", "c")
    repo.adjust_stock("WH-VEHARI", "ZINC-10", -5, "count", "owner", "owner")
    repo.set_stock("WH-VEHARI", "SEED-WHEAT", 100)
    replay = repo.replay_stock_ledger()
    for s in repo.list_stock():
        assert replay.get((s.warehouse_id, s.sku), 0) == s.on_hand, (s.warehouse_id, s.sku)


# ================================================================ gapless numbering
def test_documents_are_numbered_gaplessly_per_series(repo):
    y = _year()
    a = repo.record_payment("C-001", 100, "cash", "", "h"); b = repo.record_payment("C-002", 100, "bank", "", "h")
    c = repo.add_ledger("C-001", "credit_note", -50, "damaged bag", None, "h", "owner", "adjustment")
    ob = repo.opening_balance("C-010", 1000, "owner")
    assert [a.entry_id, b.entry_id, c.entry_id, ob.entry_id] == [f"RCP-{y}-000001", f"RCP-{y}-000002", f"CRN-{y}-000001", f"OPB-{y}-000001"]
    assert all(e.doc_no == e.entry_id for e in (a, b, c, ob))
    with pytest.raises(NotFoundError):     # a failed create takes no number
        repo.record_payment("C-999", 100, "cash", "", "h")
    assert repo.record_payment("C-001", 1, "cash", "", "h").entry_id == f"RCP-{y}-000003"


def test_a_number_taken_inside_a_failed_transaction_is_not_burnt(repo):
    from munshi.domain.models import now_iso
    from munshi.domain.repository.guarded import immediate_tx
    from munshi.domain.repository.numbering import next_doc_no
    with pytest.raises(RuntimeError):
        with immediate_tx(repo) as c:
            assert next_doc_no(repo, c, "receipt", now_iso()).endswith("-000001")
            raise RuntimeError("the insert failed after the number was taken")
    assert repo.record_payment("C-001", 1, "cash", "", "h").entry_id.endswith("-000001")
    with pytest.raises(RuntimeError, match="inside the transaction"):
        next_doc_no(repo, repo._conn.cursor(), "receipt", now_iso())


def test_numbers_restart_each_business_year_and_prefix_is_configurable(repo):
    from munshi.domain.models import now_iso
    from munshi.domain.repository.guarded import immediate_tx
    from munshi.domain.repository.numbering import next_doc_no
    repo.set_setting("invoice_prefix", "SI")
    with immediate_tx(repo) as c:
        assert next_doc_no(repo, c, "invoice", "2026-12-31T18:59:59+00:00") == "SI-2026-000001"    # 23:59:59 PKT
        assert next_doc_no(repo, c, "invoice", "2026-12-31T19:00:00+00:00") == "SI-2027-000001"    # 00:00 PKT new year
        assert next_doc_no(repo, c, "invoice", now_iso()).startswith(f"SI-{to_business_date(now_iso()).year}-")
    repo.set_setting("invoice_prefix", "RCP")                                       # would collide with receipts
    with pytest.raises(StateError, match="prefix"):
        with immediate_tx(repo) as c:
            next_doc_no(repo, c, "invoice", now_iso())


def test_invoice_document_prints_the_gapless_number_and_the_disclaimer(repo):
    from munshi.documents.invoice import TAX_DISCLAIMER, invoice_data, render_html, render_pdf, statement_html, whatsapp_text
    o = repo.create_order("C-002", [{"sku": "UREA-50", "qty": 3}], "t", "", "o")
    repo.confirm_order(o.order_id, "o", "c"); repo.allocate_order(o.order_id, "WH-MULTAN", "g")
    p = repo.create_dispatch_plan(today_iso(), "R-MULTAN-N", "V-01", [o.order_id], "g"); repo.approve_dispatch_plan(p.plan_id, "g", "c")
    st = repo.list_stops(p.plan_id)[0]
    r = repo.close_stop(st.stop_id, [{"sku": "UREA-50", "qty": 3}], [], 1000.5, st.otp, "d")
    assert r["invoice_id"] == f"INV-{_year()}-000001" and r["receipt_id"] == f"RCP-{_year()}-000001"
    d = invoice_data(repo, r["invoice_id"])
    assert d["number"] == r["invoice_id"] and d["lines"] == [{"sku": "UREA-50", "qty": 3, "unit_price": 3850.0, "total": 11550.0, "name": "Urea 50kg", "unit": "bag"}]
    assert d["paid_on_delivery"] == 1000.5
    html = render_html(d)
    assert r["invoice_id"] in html and TAX_DISCLAIMER in html and render_pdf(d)[:4] == b"%PDF"
    rev = repo.reverse_ledger_entry(r["invoice_id"], "wrong customer", "owner", "owner")
    dr = invoice_data(repo, rev.entry_id)
    assert dr["kind"] == "Reversal" and dr["lines"] == [] and "reverses" in whatsapp_text(dr)
    assert "reversal of " + r["invoice_id"] in statement_html(repo, "C-002")


# ================================================================ margin snapshot
def test_margin_of_a_past_sale_never_moves_with_later_costs(repo):
    o = repo.create_order("C-002", [{"sku": "UREA-50", "qty": 30}], "t", "", "o")
    repo.confirm_order(o.order_id, "o", "c"); repo.allocate_order(o.order_id, "WH-MULTAN", "g")
    p = repo.create_dispatch_plan(today_iso(), "R-MULTAN-N", "V-01", [o.order_id], "g"); repo.approve_dispatch_plan(p.plan_id, "g", "c")
    st = repo.list_stops(p.plan_id)[0]
    repo.close_stop(st.stop_id, [{"sku": "UREA-50", "qty": 30}], [], 0, st.otp, "d")
    t = today_iso()
    first = repo.sales_report(t, t)
    assert first["gross_margin"] == 30 * (3850 - 3600) == 7500                                  # the reviewers' example shape
    repo.record_purchase("S-001", "WH-MULTAN", [{"sku": "UREA-50", "qty": 100, "unit_cost": 3800}], "", 0, "k")
    prod = repo.get_product("UREA-50"); prod.cost_price = 9999
    repo.upsert_product(prod)
    assert repo.sales_report(t, t) == first and repo.profit_summary(t, t)["gross_margin"] == 7500
    val = repo.stock_valuation()
    assert val["costing"] == "moving_average" and 9999 not in {r["unit_cost"] for r in val["rows"]}


# ================================================================ V5 migration of a pre-V5 file
def _legacy_file(path, bad_amount=None) -> None:
    """A business file as the pre-V5 code wrote it: REAL floats, float noise, set_stock without moves,
    a negative stock row (allow_negative), a delivered stop, JSON money."""
    conn = sqlite3.connect(path, isolation_level=None)
    conn.execute("CREATE TABLE schema_version (version INTEGER PRIMARY KEY, applied_at TEXT)")
    for v, sql in ((1, migrations.V1), (2, migrations.V2), (3, migrations.V3), (4, migrations.V4)):
        for s in [x.strip() for x in sql.split(";") if x.strip()]:
            conn.execute(s)
        conn.execute("INSERT INTO schema_version VALUES (?, datetime('now'))", (v,))
    ins = conn.execute
    ins("INSERT INTO customers VALUES ('C-1','Malik','0300','standard',400000.0,NULL,'ur-en','',2.5,30,1)")
    ins("INSERT INTO products VALUES ('UREA','Urea',3850.0,'[]',1,3666.6666666666665,'bag','',10,1)")
    ins("INSERT INTO products VALUES ('DAP','DAP',6250.0,'[]',1,5900.0,'bag','',10,1)")
    ins("INSERT INTO warehouses VALUES ('WH','Godown')")
    ins("INSERT INTO stock VALUES ('WH','UREA',90,0)")          # set_stock'd: no moves at all
    ins("INSERT INTO stock VALUES ('WH','DAP',-3,0)")           # went negative through allow_negative
    ins("INSERT INTO stock_moves VALUES ('MOV-1','WH','DAP',-13,'sale','DSP-1','2026-08-01T10:00:00+00:00')")
    ins("INSERT INTO orders VALUES ('ORD-1','C-1',?, 'delivered','app','', '2026-08-01T09:00:00+00:00','WH',2.5,'','')",
        (json.dumps([{"sku": "DAP", "qty": 10, "unit_price": 6093.75}, {"sku": "UREA", "qty": 3, "unit_price": 3753.7500000000005}]),))
    ins("INSERT INTO dispatch_plans VALUES ('DSP-1','2026-08-01','R','V','WH','[\"ORD-1\"]','approved',13,'2026-08-01T09:30:00+00:00')")
    ins("INSERT INTO stops VALUES ('STP-1','DSP-1','ORD-1','C-1',1,'short',?,?,1000.1,'1234',1,'2026-08-01T12:00:00+00:00','')",
        (json.dumps([{"sku": "DAP", "qty": 10}, {"sku": "UREA", "qty": 2}]), json.dumps([])))
    ins("INSERT INTO stop_closes VALUES ('STP-1',NULL,'h','{}','2026-08-01T12:00:00+00:00')")
    ins("INSERT INTO deposits VALUES ('DEP-1','DSP-1',1000.1,'cashier','2026-08-01T18:00:00+00:00')")
    ledger = [("INV-A1B2C3D4", "invoice", 0.1 + 0.2, "ORD-1"), ("PAY-1", "payment", -0.1, "STP-1"), ("PAY-2", "payment", -0.2, "x"),
              ("INV-OB", "invoice", 68438.52, "opening balance"), ("CRN-1", "credit_note", -5.555, "x")]
    for eid, kind, amt, ref in ledger:
        ins("INSERT INTO ledger VALUES (?,?,?,?,?,NULL,'2026-08-01T12:00:00+00:00','cash','office')", (eid, "C-1", kind, bad_amount if (bad_amount and eid == "PAY-2") else amt, ref))
    ins("INSERT INTO suppliers VALUES ('S-1','Fauji','','',1)")
    ins("INSERT INTO purchases VALUES ('PUR-1','S-1','WH',?,1100110.0,'',300000.0,'2026-07-30T10:00:00+00:00')",
        (json.dumps([{"sku": "UREA", "qty": 300, "unit_cost": 3667.0333333333333}]),))
    ins("INSERT INTO supplier_ledger VALUES ('BIL-1','S-1','bill',1100110.0,'PUR-1','','2026-07-30T10:00:00+00:00')")
    ins("INSERT INTO supplier_ledger VALUES ('SPY-1','S-1','payment',-300000.0,'PUR-1','cash','2026-07-30T10:00:00+00:00')")
    ins("INSERT INTO expenses VALUES ('EXP-1','fuel',8500.499,'diesel','cash','Bilal','2026-08-01','2026-08-01T09:00:00+00:00')")
    ins("INSERT INTO reminders VALUES ('REM-1','C-1','gentle',68438.52,3,'m','drafted','2026-08-01T09:00:00+00:00')")
    ins("INSERT INTO promises VALUES ('PRM-1','C-1',5000.0,'2026-08-05','2026-08-01T09:00:00+00:00')")
    conn.close()


def test_v5_migrates_a_legacy_float_file_exactly(tmp_path):
    path = str(tmp_path / "legacy.db"); _legacy_file(path)
    repo = MunshiRepository(path)
    one = lambda q, a=(): repo._one(q, a)[0]   # noqa: E731
    assert one("SELECT MAX(version) FROM schema_version") == 5
    amounts = {r["entry_id"]: r["amount"] for r in repo._all("SELECT entry_id, amount FROM ledger")}
    assert amounts == {"INV-A1B2C3D4": 30, "PAY-1": -10, "PAY-2": -20, "INV-OB": 6843852, "CRN-1": -556}
    assert repo.outstanding_paisa("C-1") == 30 - 10 - 20 + 6843852 - 556 and repo.get_ledger_entry("INV-A1B2C3D4").doc_no is None
    assert repo.get_customer("C-1").credit_limit == 400000.0 and repo.get_product("UREA").cost_price == 3666.67
    assert repo.get_order("ORD-1").items[1].unit_price == 3753.75 and repo.get_order("ORD-1").total == 10 * 6093.75 + 3 * 3753.75
    pur = repo.get_purchase("PUR-1")
    assert pur.total == 1100110.0 and pur.items[0]["unit_cost"] == 3667.03 and repo.supplier_balance("S-1") == 800110.0
    assert repo.get_stop("STP-1").cash_collected == 1000.1 and repo.deposits("DSP-1")[0].amount_counted == 1000.1
    assert repo.expenses_between("2026-08-01", "2026-08-01")[0].amount == 8500.5          # 8500.499 -> 850050 paisa
    assert repo.get_reminder("REM-1").amount_due == 68438.52 and repo.list_promises("C-1")[0].amount == 5000.0
    # the stock ledger now replays to the stock table, legacy gaps closed by 'opening' moves
    assert repo.replay_stock_ledger() == {("WH", "UREA"): 90, ("WH", "DAP"): -3}
    assert {m.kind for m in repo.stock_moves()} == {"sale", "opening"}
    # opening inventory value = on_hand x cost price at migration; a negative holding carries no value
    assert one("SELECT value_paisa FROM inventory_value WHERE sku='UREA'") == 90 * 366667 and one("SELECT value_paisa FROM inventory_value WHERE sku='DAP'") == 0
    # the legacy delivery got a sale record, costed at the legacy cost price and flagged as such
    sl = [dict(r) for r in repo._all("SELECT sku, qty, revenue_paisa, cost_paisa, cost_basis, invoice_id FROM sale_lines ORDER BY sku")]
    assert sl == [{"sku": "DAP", "qty": 10, "revenue_paisa": 6093750, "cost_paisa": 5900000, "cost_basis": "legacy_cost_price", "invoice_id": "INV-A1B2C3D4"},
                  {"sku": "UREA", "qty": 2, "revenue_paisa": 750750, "cost_paisa": 733334, "cost_basis": "legacy_cost_price", "invoice_id": "INV-A1B2C3D4"}]
    # the legacy negative row may only move up; a further decrease is refused
    repo.move_stock("WH", "DAP", 1, "adjust", "count")
    assert repo.get_stock("WH", "DAP").on_hand == -2
    with pytest.raises(InsufficientStockError):
        repo.move_stock("WH", "DAP", -1, "adjust", "count")
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        with repo._tx() as c:
            c.execute("UPDATE ledger SET amount=0 WHERE entry_id='PAY-1'")
    assert repo._all("PRAGMA foreign_key_check") == []
    assert repo._one("SELECT name FROM sqlite_master WHERE name LIKE '%__v5'") is None
    repo.close()
    again = MunshiRepository(path)                     # re-opening applies nothing and changes nothing
    assert migrations.current_version(again._conn) == 5 and again.outstanding_paisa("C-1") == 30 - 10 - 20 + 6843852 - 556


def test_v5_cannot_half_apply(tmp_path):
    path = str(tmp_path / "bad.db"); _legacy_file(path, bad_amount="twelve rupees")
    with pytest.raises(migrations.MigrationError, match="ledger.amount of PAY-2"):
        MunshiRepository(path)
    conn = sqlite3.connect(path)
    assert conn.execute("SELECT MAX(version) FROM schema_version").fetchone()[0] == 4
    assert {r[1]: r[2] for r in conn.execute("PRAGMA table_info(customers)")}["credit_limit"] == "REAL"   # rebuilt tables rolled back too
    assert {r[1]: r[2] for r in conn.execute("PRAGMA table_info(ledger)")}["amount"] == "REAL"
    assert conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE name LIKE '%__v5' OR name IN ('sale_lines','document_counters','inventory_value')").fetchone()[0] == 0
    assert conn.execute("SELECT amount FROM ledger WHERE entry_id='INV-OB'").fetchone()[0] == 68438.52
    conn.close()


def test_excel_import_refuses_negative_stock_and_books_supplier_balance_in_paisa(repo):
    from openpyxl import Workbook

    from munshi.documents.excel import import_xlsx
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Stock"); ws.append(["warehouse_id", "sku", "on_hand"]); ws.append(["WH-MULTAN", "UREA-50", -5]); ws.append(["WH-MULTAN", "DAP-50", 7])
    ws = wb.create_sheet("Suppliers"); ws.append(["name", "phone", "address", "opening_balance"]); ws.append(["New Depot", "", "", 12345.67])
    import io
    buf = io.BytesIO(); wb.save(buf)
    r = import_xlsx(repo, buf.getvalue(), "owner")
    assert r["stock"] == 1 and any("Stock row 2" in e for e in r["errors"])
    assert repo.get_stock("WH-MULTAN", "DAP-50").on_hand == 7 and repo.get_stock("WH-MULTAN", "UREA-50").on_hand == 420
    sid = repo.find_supplier("New Depot").supplier_id
    assert repo.supplier_balance(sid) == 12345.67 and repo.supplier_balance_paisa(sid) == 1234567


def test_product_price_edit_rejects_negative(repo):
    with pytest.raises(ValueError):
        repo.upsert_product(Product("X-1", "X", -1.0))
