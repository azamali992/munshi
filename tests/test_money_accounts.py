"""Money accounts and the small journal (Stream B): accounts and routes, transfers, cash counts, bank reconciliation,
capital / drawings, loans, fixed assets and depreciation, general entries and reversals, opening balances; the cash.py
extensions (account_id on every writer, purchases by method, system categories refused, payroll rows not reversible,
payroll LOSS recoveries counted); one audit row per action; every report in the shared table shape."""
from __future__ import annotations

import json

import pytest

from munshi.domain import accounts as A
from munshi.domain.accounts import table
from munshi.domain.models import to_paisa
from munshi.domain.repository import MunshiRepository, NotFoundError, StateError
from tests.test_ledger_projection import _biz, _net
from tests.test_reconciliation_month import _load, clock  # noqa: F401  (clock is the fixture)

TABLE_KEYS = set(table("t", [], []))


def _repo(clock=None) -> tuple[MunshiRepository, str, str]:  # noqa: F811
    r = _biz(MunshiRepository())
    hbl = r.add_money_account("bank", "HBL current", "HBL", "1234567", actor="owner", approved_by="owner")["account_id"]
    jazz = r.add_money_account("wallet", "JazzCash", "JazzCash", actor="owner", approved_by="owner")["account_id"]
    return r, hbl, jazz


def _actions(r, n=50) -> list[str]:
    return [a["action"] for a in reversed(r.audit_log(n))]


def test_accounts_defaults_routes_and_closing():
    r, hbl, jazz = _repo()
    assert (hbl, jazz) == ("ACC-HBL", "ACC-JAZZCASH")
    rows = {a["account_id"]: a for a in r.list_money_accounts()["accounts"]}
    assert rows[hbl]["number_last4"] == "4567" and rows[hbl]["is_default"] and rows[jazz]["is_default"] and rows["CASH"]["is_default"]
    second = r.add_money_account("bank", "Meezan", "HBL", actor="owner", approved_by="owner")
    assert second["account_id"] == "ACC-HBL-2" and not second["is_default"]
    r.update_money_account(second["account_id"], {"is_default": True}, "owner", "owner")
    assert not r._account_row(hbl)["is_default"]
    r.set_method_route("bank", hbl, r.digest()["date"], "owner", "owner")
    with pytest.raises(StateError, match="route them elsewhere"):
        r.update_money_account(hbl, {"active": False}, "owner", "owner")
    with pytest.raises(StateError, match="cannot be closed"):
        r.update_money_account("CASH", {"active": False}, "owner", "owner")
    with pytest.raises(ValueError, match="cannot change kind"):
        r.update_money_account(jazz, {"kind": "bank"}, "owner", "owner")
    with pytest.raises(ValueError):
        r.set_method_route("cash", hbl, "2026-09-01", "owner", "owner")      # cash lands in a cash account
    r.update_money_account(jazz, {"active": False}, "owner", "owner")
    assert jazz not in {a["account_id"] for a in r.list_money_accounts()["accounts"]}
    with pytest.raises(StateError, match="closed"):
        r.transfer("CASH", jazz, 10)
    assert {"add_money_account", "set_method_route", "update_money_account"} <= set(_actions(r))


def test_opening_balance_on_an_account_is_a_journal_entry():
    r = _biz(MunshiRepository())
    a = r.add_money_account("bank", "UBL", "UBL", opening_balance=250_000, opening_date="2026-09-01", actor="owner", approved_by="owner")
    je = r.journal_entry(a["opening_journal"])
    assert je["kind"] == "opening" and [(x["code"], x["debit"], x["credit"]) for x in je["lines"]] == [("1000", 250_000.0, 0.0), ("3010", 0.0, 250_000.0)]
    assert r.account_balance_paisa(a["account_id"]) == 25_000_000
    assert _actions(r).count("add_money_account") == 1 and "journal_posted" not in _actions(r)


def test_transfers_move_money_and_reverse_once():
    r, hbl, jazz = _repo()
    r.record_capital(100_000, "bank", hbl, actor="owner", approved_by="owner")
    t = r.transfer(hbl, jazz, 40_000, ref="IBFT", actor="clerk")
    assert t["transfer_id"].startswith("XFR-") and (r.account_balance_paisa(hbl), r.account_balance_paisa(jazz)) == (6_000_000, 4_000_000)
    with pytest.raises(StateError, match="approved by the owner"):
        r.reverse_transfer(t["transfer_id"], "wrong wallet", "owner", None)
    rev = r.reverse_transfer(t["transfer_id"], "wrong wallet", "owner", "owner")
    assert rev["amount"] == -40_000.0 and r.account_balance_paisa(jazz) == 0
    with pytest.raises(StateError, match="already reversed"):
        r.reverse_transfer(t["transfer_id"], "again", "owner", "owner")
    with pytest.raises(ValueError):
        r.transfer(hbl, hbl, 5)
    book = r.account_book(jazz, "2026-01-01", "2026-12-31")
    assert (book["money_in"], book["money_out"], book["closing"]) == (40_000.0, 40_000.0, 0.0)
    assert {"transfer_between_accounts", "reverse_account_transfer", "record_capital"} <= set(_actions(r))
    assert r.cash_flow("2026-01-01", "2026-12-31")["transfers_memo"] == 0.0         # transferred and reversed


def test_cash_count_and_its_difference():
    r, _, _ = _repo()
    r.record_payment("C1", 1_000, "cash", "", "h")
    cc = r.count_cash("CASH", 950, "evening count", actor="clerk")
    assert (cc["book"], cc["counted"], cc["difference"], cc["posted_journal"]) == (1_000.0, 950.0, -50.0, None)
    posted = r.post_cash_difference(cc["count_id"], "owner", "owner")
    assert posted["posted_journal"] and r.account_balance_paisa("CASH") == 95_000
    with pytest.raises(StateError, match="already posted"):
        r.post_cash_difference(cc["count_id"], "owner", "owner")
    even = r.count_cash("CASH", 950, actor="clerk")
    with pytest.raises(StateError, match="nothing to post"):
        r.post_cash_difference(even["count_id"], "owner", "owner")
    over = r.count_cash("CASH", 1_000, actor="clerk")
    r.post_cash_difference(over["count_id"], "owner", "owner")
    today = r.digest()["date"]
    assert r.income_statement(today, today)["cash_over_short"] == 0.0     # short 50 then over 50
    assert r.cashbook(today)["net"] == 1_000.0 - 50.0 + 50.0


def test_bank_reconciliation_ticks_and_difference():
    r, hbl, _ = _repo()
    today = r.digest()["date"]
    r.set_method_route("bank", hbl, today, "owner", "owner")
    p1 = r.record_payment("C1", 5_000, "bank", "IBFT 1", "h")
    p2 = r.record_payment("C1", 3_000, "bank", "IBFT 2", "h")
    s = r.pay_supplier("S1", 2_000, "bank", "chq 7", "k", "owner")
    rec = r.reconciliation(hbl, today)
    assert (rec["book"], rec["uncleared_in_total"], rec["uncleared_out_total"], rec["statement"]) == (6_000.0, 8_000.0, 2_000.0, None)
    r.mark_cleared(hbl, [{"source": "ledger", "source_id": p1.entry_id}, {"source": "supplier_ledger", "source_id": s.entry_id}], today, True, "clerk")
    saved = r.save_reconciliation(hbl, today, 3_000, "clerk")               # the bank shows +5,000 - 2,000
    assert (saved["adjusted"], saved["difference"]) == (3_000.0, 0.0) and [x["source_id"] for x in saved["uncleared_in"]] == [p2.entry_id]
    assert r.reconciliation(hbl, today)["difference"] == 0.0
    r.mark_cleared(hbl, [{"source": "ledger", "source_id": p1.entry_id}], today, False, "clerk")   # untick
    assert r.reconciliation(hbl, today)["difference"] == 5_000.0          # the statement shows 5,000 the books now call uncleared
    with pytest.raises(NotFoundError):
        r.mark_cleared(hbl, [{"source": "ledger", "source_id": "RCP-NOPE"}], today, True, "clerk")
    acct = {a["account_id"]: a for a in r.list_money_accounts()["accounts"]}[hbl]
    assert (acct["last_reconciled"], acct["last_statement_balance"], acct["difference"]) == (today, 3_000.0, 0.0)
    assert _actions(r).count("mark_cleared") == 2 and "save_reconciliation" in _actions(r)


def test_capital_drawings_and_loans():
    r, hbl, _ = _repo()
    today = r.digest()["date"]
    r.set_method_route("bank", hbl, today, "owner", "owner")
    r.record_capital(500_000, "bank", actor="owner", approved_by="owner")
    r.record_drawing(20_000, "cash", actor="owner", approved_by="owner")
    with pytest.raises(StateError, match="no money account takes jazzcash"):
        r.record_capital(1, "jazzcash", actor="owner", approved_by="owner")
    loan = r.add_loan("Haji Rasheed", "informal", 200_000, "bank", terms="return by Eid", actor="owner", approved_by="owner")
    assert loan["loan_id"].startswith("LN-") and loan["outstanding"] == 200_000.0
    with pytest.raises(StateError, match="outstanding"):
        r.repay_loan(loan["loan_id"], 250_000, actor="owner", approved_by="owner")
    rep = r.repay_loan(loan["loan_id"], 50_000, 1_500, "bank", actor="owner", approved_by="owner")
    assert rep["outstanding"] == 150_000.0
    rpt = r.loans_report()
    assert rpt["loans"][0] | {} == rpt["loans"][0] and (rpt["received"], rpt["repaid"], rpt["interest_paid"], rpt["outstanding"]) == (200_000.0, 50_000.0, 1_500.0, 150_000.0)
    assert r.account_balance_paisa(hbl) == to_paisa(500_000 + 200_000 - 51_500)
    bs = r.balance_sheet()
    assert bs["equity"]["capital"] == 500_000.0 and bs["equity"]["drawings"] == -20_000.0 and bs["liabilities"][f"loan:{loan['loan_id']}"] == 150_000.0
    loan_je = r._standing_je("loan", loan["loan_id"], kind="loan")
    with pytest.raises(StateError, match="reverse the repayment"):
        r.reverse_journal(loan_je, "keyed twice", "owner", "owner")
    cf = r.cash_flow(today, today)
    assert cf["financing"]["lines"] == {"Capital introduced": 500_000.0, "Drawings": -20_000.0, "Loans received": 200_000.0, "Loans repaid": -50_000.0}
    assert cf["operating"]["lines"] == {"Interest and bank charges paid": -1_500.0} and cf["ties"]
    assert [a for a in _actions(r) if a in A.AUDIT_ACTION] == ["add_money_account", "add_money_account", "set_method_route", "record_capital",
                                                               "record_drawing", "record_loan", "repay_loan"]


def test_fixed_assets_depreciate_once_per_month_and_dispose():
    r, hbl, _ = _repo()
    r.record_capital(3_000_000, "bank", hbl, on_date="2026-01-01", actor="owner", approved_by="owner")
    fa = r.add_fixed_asset("Shehzore LEU-1", "vehicle", 2_400_000, "2026-01-15", 60, salvage=0, method="bank", account_id=hbl,
                           vehicle_id="V1", actor="owner", approved_by="owner")
    assert fa["asset_id"].startswith("FA-") and fa["monthly_dep"] == 40_000.0 and fa["status"] == "in use"
    run = r.run_depreciation("2026-03", "owner", "owner")
    assert [x["period"] for x in run["posted"]] == ["2026-01", "2026-02", "2026-03"] and run["total"] == 120_000.0
    assert r.run_depreciation("2026-03", "owner", "owner")["posted"] == []          # idempotent per asset-month
    feb = next(x["journal"] for x in run["posted"] if x["period"] == "2026-02")
    r.reverse_journal(feb, "posted twice by hand", "owner", "owner")
    assert [x["period"] for x in r.run_depreciation("2026-03", "owner", "owner")["posted"]] == ["2026-02"]   # a reversed month comes back
    reg = r.fixed_assets_register()["assets"][0]
    assert (reg["acc_dep"], reg["book_value"], reg["depreciated_through"]) == (120_000.0, 2_280_000.0, "2026-03")
    land = r.add_fixed_asset("Godown plot", "land", 1_000_000, "2026-02-01", 0, funded_by="payable", actor="owner", approved_by="owner")
    assert r.run_depreciation("2026-04", "owner", "owner")["posted"][0]["name"] == "Shehzore LEU-1" and land["status"] == "not depreciated"
    with pytest.raises(StateError, match="first"):
        r.reverse_journal(r._standing_je("fixed_asset", fa["asset_id"], kind="asset_purchase"), "wrong", "owner", "owner")
    d = r.dispose_fixed_asset(fa["asset_id"], "2026-04-30", 2_300_000, "bank", hbl, "owner", "owner")
    assert d["status"] == "disposed" and d["gain"] == 60_000.0                      # 2,300,000 - (2,400,000 - 160,000)
    with pytest.raises(StateError, match="already disposed"):
        r.dispose_fixed_asset(fa["asset_id"], "2026-05-01", 1, "bank", hbl, "owner", "owner")
    assert r.run_depreciation("2026-06", "owner", "owner")["posted"] == []         # nothing after the disposal month
    assert r.balance_sheet()["balanced"] and _net(r, A.FIXED_ASSETS_COST) == 100_000_000
    inc = r.income_statement("2026-01-01", "2026-12-31")                  # the February reversal is dated today
    assert inc["depreciation"] == 160_000.0 and inc["other_income"] == 60_000.0


def test_general_entries_refuse_subledgers_and_reverse():
    r, hbl, _ = _repo()
    today = r.digest()["date"]
    with pytest.raises(ValueError, match="kept by its own records"):
        r.post_journal(today, "adjustment", "fix khata", [{"code": "1100", "debit": 5}, {"code": "8000", "credit": 5}], actor="owner", approved_by="owner")
    with pytest.raises(ValueError, match="must balance"):
        r.post_journal(today, "bank_charge", "charges", [{"code": "7000", "debit": 5}, {"code": "1000", "account_id": hbl, "credit": 4}], actor="owner", approved_by="owner")
    with pytest.raises(ValueError, match="own actions"):
        r.post_journal(today, "capital", "capital", [{"code": "1000", "account_id": hbl, "debit": 5}, {"code": "3000", "credit": 5}], actor="owner", approved_by="owner")
    je = r.post_journal(today, "bank_charge", "HBL SMS charges", [{"code": "7000", "debit": 150}, {"code": "1000", "account_id": hbl, "credit": 150}],
                        actor="owner", approved_by="owner")
    assert je["je_id"].startswith("JV-") and r.account_balance_paisa(hbl) == -15_000
    rev = r.reverse_journal(je["je_id"], "bank refunded it", "owner", "owner")
    assert rev["reversal_of"] == je["je_id"] and r.account_balance_paisa(hbl) == 0
    with pytest.raises(StateError, match="itself a reversal"):
        r.reverse_journal(rev["je_id"], "undo", "owner", "owner")
    gj = r.general_journal(today, today)
    assert gj["debits"] == gj["credits"] == 300.0 and gj["table"]["count"] == 4


def test_opening_balances_bring_money_assets_and_loans_in_one_entry():
    r, hbl, _ = _repo()
    je = r.record_opening_balances("2026-07-01", money=[{"account_id": "CASH", "amount": 50_000}, {"account_id": hbl, "amount": -10_000}],
                                   assets=[{"name": "Suzuki pickup", "category": "vehicle", "cost": 1_200_000, "life_months": 60, "acquired_on": "2025-07-01",
                                            "accumulated_depreciation": 240_000}],
                                   loans=[{"lender": "Bank Alfalah", "kind": "bank", "amount": 300_000, "terms": "KIBOR+3"}], actor="owner", approved_by="owner")
    assert je["kind"] == "opening" and sum(x["debit"] for x in je["lines"]) == sum(x["credit"] for x in je["lines"])
    eq = next(x for x in je["lines"] if x["code"] == "3010")
    assert eq["credit"] == 50_000 - 10_000 + 1_200_000 - 240_000 - 300_000
    run = r.run_depreciation("2026-07", "owner", "owner")
    assert run["posted"] == [{"asset_id": je["assets"][0], "name": "Suzuki pickup", "period": "2026-07", "amount": 20_000.0, "journal": run["posted"][0]["journal"]}]
    assert r.loans_report()["outstanding"] == 300_000.0 and r.balance_sheet()["balanced"]
    assert r.cash_flow("2026-07-01", "2026-07-31")["opening"] == 40_000.0


def test_cash_writers_take_an_account_and_old_calls_are_unchanged(clock):  # noqa: F811
    r, hbl, jazz = _repo()
    clock.at(2, 10)
    # old calls: exactly as before (cash, the drawer; no account_id stored)
    p = r.record_purchase("S1", "WH-A", [{"sku": "UREA", "qty": 5, "unit_cost": 3600}], "B-1", 10_000, "k")
    row = r._one("SELECT method, account_id FROM supplier_ledger WHERE ref=? AND kind='payment'", (p.purchase_id,))
    assert (row["method"], row["account_id"]) == ("cash", None) and r.cashbook("2026-09-02")["total_out"] == 10_000.0
    # new: paid by bank from HBL
    p2 = r.record_purchase("S1", "WH-A", [{"sku": "UREA", "qty": 5, "unit_cost": 3600}], "B-2", 18_000, "k", method="bank", account_id=hbl)
    row = r._one("SELECT method, account_id FROM supplier_ledger WHERE ref=? AND kind='payment'", (p2.purchase_id,))
    assert (row["method"], row["account_id"]) == ("bank", hbl) and _net(r, A.MONEY, hbl) == -1_800_000
    r.reverse_purchase(p2.purchase_id, "wrong bill", "owner", "owner")
    assert _net(r, A.MONEY, hbl) == 0                                     # the reversal goes back into the same account
    r.pay_supplier("S1", 1_000, "jazzcash", "", "k", "owner", account_id=jazz)
    e = r.record_expense("rent", 700, "shop", "bank", "B", "h", account_id=hbl)
    r.reverse_expense(e.expense_id, "wrong month", "owner", "owner")
    assert _net(r, A.MONEY, jazz) == -100_000 and _net(r, A.MONEY, hbl) == 0
    plan, stops = _load(r, "R-A", [("C1", [{"sku": "UREA", "qty": 1}])])
    st = stops["C1"]; r.close_stop(st.stop_id, [{"sku": "UREA", "qty": 1}], [], 4_000, st.otp, "d")
    r.record_deposit(plan.plan_id, 4_000, "cashier", "h", account_id=hbl)   # banked straight away: not in the drawer's cashbook
    assert _net(r, A.MONEY, hbl) == 400_000 and r.cashbook("2026-09-02")["total_handins"] == 0.0
    audit = [a for a in r.audit_log(30) if a["action"] == "record_purchase"][0]
    assert audit["payload"]["account"] == hbl and audit["payload"]["method"] == "bank"


def test_payroll_rows_are_payrolls_own():
    r = _biz(MunshiRepository())
    for cat in A.SYSTEM_EXPENSE_CATEGORIES:
        with pytest.raises(StateError, match="payroll"):
            r.record_expense(cat, 100, "x", "cash", "B", "h")
    assert r.record_expense("salary", 100, "old habit", "cash", "B", "h").category == "salary"      # hand-keyed 'salary' stays legal
    r._conn.execute("INSERT INTO expenses (expense_id, category, amount, note, method, paid_by, expense_date, created_at) VALUES "
                    "('EXP-P1','staff_wages',100,'PAY-2026-000003 September payroll','payroll','payroll','2026-09-30','2026-09-30T10:00:00+00:00')")
    with pytest.raises(StateError, match="reverse payroll run PAY-2026-000003"):
        r.reverse_expense("EXP-P1", "wrong", "owner", "owner")


def test_a_payroll_loss_recovery_counts_against_the_plans_shortage(clock):  # noqa: F811
    r = _biz(MunshiRepository())
    clock.at(3, 9)
    r.set_stock("WH-A", "UREA", 10)
    plan, stops = _load(r, "R-A", [("C1", [{"sku": "UREA", "qty": 1}])])
    st = stops["C1"]; r.close_stop(st.stop_id, [{"sku": "UREA", "qty": 1}], [], 4_000, st.otp, "d")
    first = r.record_deposit(plan.plan_id, 3_000, "cashier", "h")
    assert first["shortage_entry"]["amount"] == 1_000.0
    # payroll deducted 400 of it from the driver's salary: a negative cash_shortage row, method payroll (Stream A)
    r._conn.execute("INSERT INTO expenses (expense_id, category, amount, note, method, paid_by, expense_date, created_at) VALUES (?,?,?,?,?,?,?,?)",
                    ("EXP-REC1", "cash_shortage", -40_000, f"RECOVERY {plan.plan_id} PAY-2026-000001", "payroll", "payroll", "2026-09-30", "2026-09-30T10:00:00+00:00"))
    # the driver then hands in the other 600: the plan is square -- 1,000 short, 400 through payroll, 600 in cash
    last = r.record_deposit(plan.plan_id, 600, "cashier", "h")
    assert last["variance"] == -400.0 and last["shortage_entry"]["amount"] == -600.0
    assert r._one("SELECT SUM(amount) s FROM expenses WHERE category='cash_shortage'")["s"] == 0
    # and if he hands in 400 more (paid twice), nothing is "recovered" into a gain: cash with drivers flags it instead
    extra = r.record_deposit(plan.plan_id, 400, "cashier", "h")
    assert extra["shortage_entry"] is None and r._one("SELECT SUM(amount) s FROM expenses WHERE category='cash_shortage'")["s"] == 0


def test_every_report_is_in_the_shared_table_shape():
    r, hbl, _ = _repo()
    today = r.digest()["date"]
    reports = [r.list_money_accounts(), r.account_book("CASH", today, today), r.method_routes(), r.reconciliation(hbl, today), r.trial_balance(),
               r.general_journal(today, today), r.income_statement(today, today), r.balance_sheet(), r.cash_flow(today, today), r.owner_kpis(),
               r.margins("product", today, today), r.margins("route", today, today), r.fixed_assets_register(), r.loans_report(), r.verify_books(),
               r.period_status(), r.run_depreciation(today[:7], "owner", "owner")]
    for rep in reports:
        assert set(rep["table"]) == TABLE_KEYS, rep["table"]["title"]
        json.dumps(rep)                                                   # the console and chat can serialise every one
    assert len(r.income_statement(today, today)["tables"]) == 2


def test_the_export_carries_the_accountants_sheets():
    import io

    from openpyxl import load_workbook

    from munshi.documents.excel import export_xlsx
    from munshi.domain.seed import seeded_repository
    wb = load_workbook(io.BytesIO(export_xlsx(seeded_repository())))
    assert {"GeneralJournal", "TrialBalance", "Books"} <= set(wb.sheetnames)
    tb = wb["TrialBalance"]
    last = [c.value for c in tb[tb.max_row]]
    assert last[1] == "Total" and last[2] == last[3]                      # debits = credits
    assert [c.value for c in wb["GeneralJournal"][1]] == ["date", "doc", "source", "code", "account", "party", "debit", "credit", "memo"]
